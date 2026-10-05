"""Gera a planilha gratuita "Modelo integrado" (DRE, Balanço e Fluxo de Caixa ligados, 5 anos).

Abas: Comece aqui, Premissas, DRE, Balanço, Fluxo de Caixa, Ligações (as 8 conexões conferidas ao vivo),
Checagem, Lucro x Caixa (com gráfico) e Exercícios (testes guiados, inclusive quebrar uma ligação).
Convenções: Calibri, premissas em azul sobre creme, links entre abas em verde, negativos entre parênteses.

Uso: python3 scripts/gerar-planilha-modelo-integrado.py public/materiais/modelo-integrado-simplificado.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]

NAVY, INK, INK2, BLUE, GREEN, AMBER_TXT = "070D24", "0F172A", "475569", "1D4ED8", "15803D", "B45309"
INPUT_FILL, TOTAL_FILL = "FFF7E0", "F1F5F9"
NUM = '#,##0;(#,##0);"-"'
PCT = "0%"
DAYS = '0 "dias"'
URL = "https://www.mfnapratica.com.br/fundamentos?utm_source=planilha-gratuita&utm_medium=isca&utm_campaign=modelo-integrado"
# Aba com a oferta do curso: cada link com utm_content próprio para medir de onde veio a venda
CURSO = "Curso Fundamentos"
CURSO_URL = URL + "&utm_content=aba-curso"
CHECKOUT_URL = "https://pay.hotmart.com/F106435738T?src=planilha-gratuita"

YEARS = 5
COLS = ["C", "D", "E", "F", "G", "H"]  # Ano 0 .. Ano 5
PROJ = COLS[1:]


def font(color=INK, size=11, bold=False):
    return Font(name="Calibri", size=size, bold=bold, color=color)


thin = Side(style="thin", color="D9DEE7")
navy_top = Side(style="thin", color=NAVY)


def sheet(wb, name, tab, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    return ws


def title(ws, text, sub):
    ws["B2"] = text
    ws["B2"].font = font(NAVY, 18, True)
    ws["B3"] = sub
    ws["B3"].font = font(INK2, 10)


def header(ws, row, label, cols, labels):
    ws[f"B{row}"] = label
    ws[f"B{row}"].font = font("FFFFFF", 11, True)
    ws[f"B{row}"].fill = PatternFill("solid", fgColor=NAVY)
    for c, l in zip(cols, labels):
        cell = ws[f"{c}{row}"]
        cell.value = l
        cell.font = font("FFFFFF", 11, True)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="right")


def line(ws, row, label, values, kind="calc", fmt=NUM):
    """kind: calc (preto), link (verde, vem de outra aba), total (negrito com fundo)."""
    ws[f"B{row}"] = label
    total = kind == "total"
    ws[f"B{row}"].font = font(NAVY if total else INK, 11, total)
    if total:
        ws[f"B{row}"].fill = PatternFill("solid", fgColor=TOTAL_FILL)
        ws[f"B{row}"].border = Border(top=navy_top)
    for c, v in values.items():
        cell = ws[f"{c}{row}"]
        cell.value = v
        cell.number_format = fmt
        is_link = kind == "link" or (isinstance(v, str) and v.startswith("=") and _pure_link(v))
        cell.font = font(GREEN if is_link and not total else "000000", 11, total)
        if total:
            cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
            cell.border = Border(top=navy_top)


def _pure_link(f):
    # "=Aba!X1" ou "=-Aba!X1" ou "='Aba'!X1": só traz o número de outra aba
    body = f[1:].lstrip("-")
    return "!" in body and not any(op in body.split("!", 1)[1] for op in "+-*/(")


def years_header(ws, row, label, with_zero=True):
    cols = COLS if with_zero else PROJ
    labels = [f"Ano {i}" for i in (range(0, YEARS + 1) if with_zero else range(1, YEARS + 1))]
    header(ws, row, label, cols, labels)


wb = Workbook()

# ---------------------------------------------------------------- Comece aqui
start = sheet(wb, "Comece aqui", "F59E0B", {"A": 3, "B": 4, "C": 100}, first=True)

# ---------------------------------------------------------------- Premissas
pre = sheet(wb, "Premissas", BLUE, {"A": 3, "B": 44, "C": 14, "D": 46})
title(pre, "Premissas", "Mude só as células azuis. Valores em R$ mil.")


def section(ws, row, text):
    ws[f"B{row}"] = text
    ws[f"B{row}"].font = font(NAVY, 12, True)


def premissa(row, label, value, fmt, note):
    pre[f"B{row}"] = label
    pre[f"B{row}"].font = font()
    c = pre[f"C{row}"]
    c.value = value
    c.number_format = fmt
    c.font = font(BLUE)
    c.fill = PatternFill("solid", fgColor=INPUT_FILL)
    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    pre[f"D{row}"] = note
    pre[f"D{row}"].font = font(INK2, 10)


section(pre, 6, "Operação")
premissa(7, "Receita no ano 1", 1000, NUM, "Faturamento líquido do primeiro ano")
premissa(8, "Crescimento da receita por ano", 0.1, PCT, "Aplicado do ano 2 ao ano 5")
premissa(9, "Custos (% da receita)", 0.6, PCT, "Custo dos produtos ou serviços vendidos")
premissa(10, "Despesas (% da receita)", 0.2, PCT, "Vendas, gerais e administrativas")
premissa(11, "Depreciação por ano", 50, NUM, "Reduz o lucro, mas não sai do caixa")
premissa(12, "Investimento por ano (capex)", 60, NUM, "Compra de máquinas e equipamentos. Não passa pela DRE")
premissa(13, "Alíquota de IR", 0.3, PCT, "Sobre o lucro antes do imposto")
section(pre, 15, "Capital de giro")
premissa(16, "Prazo de recebimento dos clientes (dias)", 30, NUM, "Vendas que ainda não viraram caixa")
premissa(17, "Prazo de pagamento aos fornecedores (dias)", 30, NUM, "Custos que ainda não saíram do caixa")
section(pre, 19, "Saldos iniciais (Balanço do ano 0)")
premissa(20, "Caixa", 100, NUM, "")
premissa(21, "Imobilizado", 500, NUM, "")
premissa(22, "Patrimônio líquido", 600, NUM, "Ativo = Passivo + PL já no ano 0")

# ---------------------------------------------------------------- DRE
dre = sheet(wb, "DRE", NAVY, {"A": 3, "B": 34, **{c: 12 for c in COLS}})
title(dre, "DRE (Demonstração do Resultado)", "R$ mil. O lucro líquido vai para o Patrimônio líquido (Balanço) e para o Fluxo de Caixa.")
years_header(dre, 5, None)
dre["B5"] = None
dre.freeze_panes = "C6"
receita = {"D": "=Premissas!$C$7"}
for prev, c in zip(PROJ, PROJ[1:]):
    receita[c] = f"={prev}6*(1+Premissas!$C$8)"
line(dre, 6, "Receita líquida", receita)
line(dre, 7, "(−) Custos", {c: f"=-{c}6*Premissas!$C$9" for c in PROJ})
line(dre, 8, "Lucro bruto", {c: f"={c}6+{c}7" for c in PROJ}, "total")
line(dre, 9, "(−) Despesas", {c: f"=-{c}6*Premissas!$C$10" for c in PROJ})
line(dre, 10, "EBITDA", {c: f"={c}8+{c}9" for c in PROJ}, "total")
line(dre, 11, "(−) Depreciação", {c: "=-Premissas!$C$11" for c in PROJ})
line(dre, 12, "EBIT (lucro operacional)", {c: f"={c}10+{c}11" for c in PROJ}, "total")
line(dre, 13, "(−) IR", {c: f"=-MAX(0,{c}12)*Premissas!$C$13" for c in PROJ})
line(dre, 14, "Lucro líquido", {c: f"={c}12+{c}13" for c in PROJ}, "total")

# ---------------------------------------------------------------- Balanço
bal = sheet(wb, "Balanço", NAVY, {"A": 3, "B": 34, **{c: 12 for c in COLS}})
title(bal, "Balanço Patrimonial", "R$ mil. O caixa vem do Fluxo de Caixa; o patrimônio soma o lucro da DRE.")
years_header(bal, 5, "Ativo")
bal.freeze_panes = "C6"
line(bal, 6, "Caixa", {"C": "=Premissas!$C$20", **{c: f"='Fluxo de Caixa'!{c}14" for c in PROJ}})
line(bal, 7, "Contas a receber", {"C": 0, **{c: f"=DRE!{c}6*Premissas!$C$16/360" for c in PROJ}})
imob = {"C": "=Premissas!$C$21"}
for prev, c in zip(COLS, PROJ):
    imob[c] = f"={prev}8+Premissas!$C$12-Premissas!$C$11"
line(bal, 8, "Imobilizado", imob)
line(bal, 9, "Ativo total", {c: f"=SUM({c}6:{c}8)" for c in COLS}, "total")
years_header(bal, 11, "Passivo e patrimônio líquido")
line(bal, 12, "Fornecedores", {"C": 0, **{c: f"=-DRE!{c}7*Premissas!$C$17/360" for c in PROJ}})
pl = {"C": "=Premissas!$C$22"}
for prev, c in zip(COLS, PROJ):
    pl[c] = f"={prev}13+DRE!{c}14"
line(bal, 13, "Patrimônio líquido", pl)
line(bal, 14, "Passivo + PL", {c: f"={c}12+{c}13" for c in COLS}, "total")
line(bal, 16, "Checagem: Ativo − (Passivo + PL)", {c: f"=ROUND({c}9-{c}14,6)" for c in COLS})
bal["B16"].font = font(NAVY, 11, True)
for c in COLS:
    bal[f"{c}16"].font = font("000000", 11, True)
bal["B17"] = "Contas a receber e fornecedores começam em zero no ano 0 (empresa-exemplo começando a operar)."
bal["B17"].font = font(INK2, 10)

# ---------------------------------------------------------------- Fluxo de Caixa
fc = sheet(wb, "Fluxo de Caixa", NAVY, {"A": 3, "B": 40, **{c: 12 for c in COLS}})
title(fc, "Fluxo de Caixa (método indireto)", "R$ mil. Parte do lucro e explica por que o caixa não é igual ao lucro.")
years_header(fc, 5, None, with_zero=False)
fc["B5"] = None
fc.freeze_panes = "D6"
prev_of = dict(zip(PROJ, COLS))
line(fc, 6, "Lucro líquido", {c: f"=DRE!{c}14" for c in PROJ})
line(fc, 7, "(+) Depreciação", {c: f"=-DRE!{c}11" for c in PROJ})
line(fc, 8, "(−) Aumento de contas a receber", {c: f"=-(Balanço!{c}7-Balanço!{prev_of[c]}7)" for c in PROJ})
line(fc, 9, "(+) Aumento de fornecedores", {c: f"=Balanço!{c}12-Balanço!{prev_of[c]}12" for c in PROJ})
line(fc, 10, "Caixa das operações", {c: f"=SUM({c}6:{c}9)" for c in PROJ}, "total")
line(fc, 11, "(−) Investimentos (capex)", {c: "=-Premissas!$C$12" for c in PROJ})
line(fc, 12, "Variação do caixa", {c: f"={c}10+{c}11" for c in PROJ}, "total")
line(fc, 13, "Caixa inicial", {c: f"=Balanço!{prev_of[c]}6" for c in PROJ})
line(fc, 14, "Caixa final", {c: f"={c}13+{c}12" for c in PROJ}, "total")

# ---------------------------------------------------------------- Ligações
lig = sheet(wb, "Ligações", "F59E0B", {"A": 3, "B": 4, "C": 36, "D": 26, "E": 34, "F": 12, "G": 12, "H": 11, "I": 12, "J": 58})
title(lig, "Mapa das ligações", "As 8 conexões que fazem os três demonstrativos virarem um modelo só. Cada uma é conferida ao vivo.")
lig["B5"] = "Ano para conferir"
lig["B5"].font = font(NAVY, 11, True)
lig["E5"] = "Digite de 1 a 5"
lig["E5"].font = font(INK2, 10)
year = lig["D5"]
year.value = 1
year.font = font(BLUE, 11, True)
year.fill = PatternFill("solid", fgColor=INPUT_FILL)
year.border = Border(top=thin, bottom=thin, left=thin, right=thin)
year.alignment = Alignment(horizontal="center")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", showErrorMessage=True,
                    errorTitle="Ano inválido", error="Digite um ano de 1 a 5.")
lig.add_data_validation(dv)
dv.add("D5")

heads = ["#", "Ligação", "Sai de", "Chega em", "Na origem", "No destino", "Diferença", "Status", "Por que importa"]
for c, h in zip("BCDEFGHIJ", heads):
    cell = lig[f"{c}7"]
    cell.value = h
    cell.font = font("FFFFFF", 11, True)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="right" if c in "FGHI" else "left")

Y = "$D$5"


def cur(sheet_ref, row):  # valor do ano escolhido
    return f"INDEX({sheet_ref}!$D${row}:$H${row},{Y})"


def prev(sheet_ref, row):  # valor do ano anterior (coluna C é o ano 0)
    return f"INDEX({sheet_ref}!$C${row}:$G${row},{Y})"


FCR = "'Fluxo de Caixa'"
LINKS = [
    ("O lucro abre o Fluxo de Caixa", "DRE · Lucro líquido", "Fluxo de Caixa · Lucro líquido",
     f"={cur('DRE', 14)}", f"={cur(FCR, 6)}",
     "O Fluxo de Caixa indireto começa no lucro e explica, linha a linha, por que o caixa é diferente."),
    ("O lucro soma no patrimônio líquido", "DRE · Lucro líquido", "Balanço · aumento do PL",
     f"={cur('DRE', 14)}", f"={cur('Balanço', 13)}-{prev('Balanço', 13)}",
     "Sem ela, o Balanço fica com uma diferença exatamente igual ao lucro do ano."),
    ("A depreciação volta somando no caixa", "DRE · Depreciação", "Fluxo de Caixa · (+) Depreciação",
     f"=-{cur('DRE', 11)}", f"={cur(FCR, 7)}",
     "Ela reduz o lucro, mas não sai do banco. Por isso é somada de volta."),
    ("A depreciação reduz o imobilizado", "DRE · Depreciação", "Balanço · Imobilizado",
     f"=-{cur('DRE', 11)}", f"={prev('Balanço', 8)}-{cur(FCR, 11)}-{cur('Balanço', 8)}",
     "O desgaste do ano sai do valor das máquinas no Balanço."),
    ("O investimento aumenta o imobilizado", "Fluxo de Caixa · (−) Investimentos", "Balanço · Imobilizado",
     f"=-{cur(FCR, 11)}", f"={cur('Balanço', 8)}-{prev('Balanço', 8)}-{cur('DRE', 11)}",
     "Comprar máquina tira caixa, mas não passa pela DRE: vira ativo."),
    ("Contas a receber consome caixa", "Balanço · Contas a receber", "Fluxo de Caixa · (−) Contas a receber",
     f"=-({cur('Balanço', 7)}-{prev('Balanço', 7)})", f"={cur(FCR, 8)}",
     "Venda a prazo entra no lucro, mas o dinheiro ainda está com o cliente."),
    ("Fornecedores seguram caixa", "Balanço · Fornecedores", "Fluxo de Caixa · (+) Fornecedores",
     f"={cur('Balanço', 12)}-{prev('Balanço', 12)}", f"={cur(FCR, 9)}",
     "Comprar a prazo entra no custo, mas o dinheiro só sai depois."),
    ("O caixa final fecha o Balanço", "Fluxo de Caixa · Caixa final", "Balanço · Caixa",
     f"={cur(FCR, 14)}", f"={cur('Balanço', 6)}",
     "É a última ligação: com ela, Ativo = Passivo + PL."),
]
for i, (name, src, dst, fo, fd, why) in enumerate(LINKS):
    r = 8 + i
    lig[f"B{r}"] = i + 1
    lig[f"C{r}"] = name
    lig[f"D{r}"] = src
    lig[f"E{r}"] = dst
    lig[f"F{r}"] = fo
    lig[f"G{r}"] = fd
    lig[f"H{r}"] = f"=ROUND(F{r}-G{r},6)"
    lig[f"I{r}"] = f'=IF(ABS(H{r})<0.01,"Ligada","Quebrada")'
    lig[f"J{r}"] = why
    lig[f"B{r}"].font = font(INK2)
    lig[f"C{r}"].font = font(INK, 11, True)
    for c in "DE":
        lig[f"{c}{r}"].font = font(INK2, 10)
    for c in "FG":
        lig[f"{c}{r}"].font = font(GREEN)
        lig[f"{c}{r}"].number_format = NUM
    lig[f"H{r}"].number_format = NUM
    lig[f"H{r}"].font = font("000000")
    lig[f"I{r}"].font = font("000000", 11, True)
    lig[f"I{r}"].alignment = Alignment(horizontal="right")
    lig[f"J{r}"].font = font(INK2, 10)
    lig[f"J{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    for c in "BCDEFGHIJ":
        lig[f"{c}{r}"].border = Border(bottom=thin)
        if c != "J":
            lig[f"{c}{r}"].alignment = Alignment(vertical="center", horizontal=lig[f"{c}{r}"].alignment.horizontal, wrap_text=c in "CDE")
    lig.row_dimensions[r].height = 30
last = 8 + len(LINKS) - 1
lig[f"C{last + 2}"] = "Ligações funcionando"
lig[f"C{last + 2}"].font = font(NAVY, 13, True)
lig[f"D{last + 2}"] = f'=COUNTIF(I8:I{last},"Ligada")&" de {len(LINKS)}"'
lig[f"D{last + 2}"].font = font(NAVY, 13, True)
lig[f"C{last + 3}"] = "Para ver uma ligação quebrar, faça o exercício 5 da aba Exercícios."
lig[f"C{last + 3}"].font = font(INK2, 10)
lig.freeze_panes = "B8"

# ---------------------------------------------------------------- Checagem
chk = sheet(wb, "Checagem", GREEN, {"A": 3, "B": 44, **{c: 12 for c in COLS}})
title(chk, "Checagem", "Tudo zero = o modelo fecha. Qualquer número diferente de zero aponta uma ligação quebrada.")
years_header(chk, 5, None)
chk["B5"] = None
line(chk, 6, "Ativo − (Passivo + PL)", {c: f"=Balanço!{c}16" for c in COLS})
line(chk, 7, "Caixa do Fluxo de Caixa − Caixa do Balanço", {c: f"=ROUND('Fluxo de Caixa'!{c}14-Balanço!{c}6,6)" for c in PROJ})
line(chk, 8, "Lucro da DRE − aumento do PL", {c: f"=ROUND(DRE!{c}14-(Balanço!{c}13-Balanço!{prev_of[c]}13),6)" for c in PROJ})
chk["B9"] = "Situação"
chk["B9"].font = font(NAVY, 11, True)
for c in COLS:
    cell = chk[f"{c}9"]
    cell.value = f'=IF(AND(ABS({c}6)<0.01,ABS(N({c}7))<0.01,ABS(N({c}8))<0.01),"Fecha","Não fecha")'
    cell.font = font("000000", 11, True)
    cell.alignment = Alignment(horizontal="right")
chk["B11"] = "Resultado do modelo"
chk["B11"].font = font(NAVY, 13, True)
chk["C11"] = '=IF(COUNTIF(C9:H9,"Fecha")=6,"✓ O modelo fecha nos 5 anos","✗ Há uma ligação quebrada")'
chk["C11"].font = font(NAVY, 13, True)
chk["C11"].alignment = Alignment(horizontal="center")
chk.merge_cells("C11:H11")

# ---------------------------------------------------------------- Lucro x Caixa
lc = sheet(wb, "Lucro x Caixa", NAVY, {"A": 3, "B": 40, **{c: 12 for c in COLS}})
title(lc, "Lucro x Caixa", "R$ mil. Por que o caixa das operações não é igual ao lucro, ano a ano.")
years_header(lc, 5, None, with_zero=False)
lc["B5"] = None
line(lc, 6, "Lucro líquido", {c: f"=DRE!{c}14" for c in PROJ})
line(lc, 7, "(+) Depreciação", {c: f"='Fluxo de Caixa'!{c}7" for c in PROJ})
line(lc, 8, "(−) Aumento de contas a receber", {c: f"='Fluxo de Caixa'!{c}8" for c in PROJ})
line(lc, 9, "(+) Aumento de fornecedores", {c: f"='Fluxo de Caixa'!{c}9" for c in PROJ})
line(lc, 10, "Caixa das operações", {c: f"=SUM({c}6:{c}9)" for c in PROJ}, "total")
line(lc, 11, "Diferença: caixa − lucro", {c: f"={c}10-{c}6" for c in PROJ})
line(lc, 12, "Conversão do lucro em caixa", {c: f"=IF({c}6=0,0,{c}10/{c}6)" for c in PROJ}, fmt=PCT)
years_header(lc, 14, "Indicadores", with_zero=False)
line(lc, 15, "Margem EBITDA", {c: f"=IF(DRE!{c}6=0,0,DRE!{c}10/DRE!{c}6)" for c in PROJ}, fmt=PCT)
line(lc, 16, "Margem líquida", {c: f"=IF(DRE!{c}6=0,0,DRE!{c}14/DRE!{c}6)" for c in PROJ}, fmt=PCT)
line(lc, 17, "Prazo médio de recebimento", {c: f"=IF(DRE!{c}6=0,0,Balanço!{c}7/DRE!{c}6*360)" for c in PROJ}, fmt=DAYS)
line(lc, 18, "Prazo médio de pagamento", {c: f"=IF(DRE!{c}7=0,0,Balanço!{c}12/-DRE!{c}7*360)" for c in PROJ}, fmt=DAYS)
line(lc, 19, "Caixa no fim do ano", {c: f"=Balanço!{c}6" for c in PROJ})
lc["B20"] = "Conversão acima de 100%: a operação gera mais caixa do que lucro. Abaixo: o lucro está preso no capital de giro."
lc["B20"].font = font(INK2, 10)

chart = BarChart()
chart.type = "col"
chart.grouping = "clustered"
chart.title = "Lucro líquido x Caixa das operações (R$ mil)"
chart.y_axis.numFmt = "#,##0"
chart.y_axis.majorGridlines = None
chart.height = 8
chart.width = 18
cats = Reference(lc, min_col=4, max_col=8, min_row=5)
for row, color in ((6, NAVY), (10, "F59E0B")):
    data = Reference(lc, min_col=2, max_col=8, min_row=row)
    chart.add_data(data, from_rows=True, titles_from_data=True)
    s = chart.series[-1]
    s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.solidFill = color
chart.set_categories(cats)
chart.legend.position = "b"
lc.add_chart(chart, "B22")
# A série precisa começar na coluna D (ano 1): ajusta a referência dos valores
for s, row in zip(chart.series, (6, 10)):
    s.val.numRef.f = f"'Lucro x Caixa'!$D${row}:$H${row}"

# ---------------------------------------------------------------- Exercícios
ex = sheet(wb, "Exercícios", "F59E0B", {"A": 3, "B": 4, "C": 30, "D": 46, "E": 52, "F": 13, "G": 13, "H": 13})
title(ex, "Exercícios: teste as ligações", "Faça um de cada vez e volte o valor original antes do próximo. Os números à direita mudam na hora.")
heads = ["#", "Experimento", "O que fazer", "O que observar", "Lucro ano 1", "Caixa operações ano 1", "Caixa final ano 5"]
for c, h in zip("BCDEFGH", heads):
    cell = ex[f"{c}5"]
    cell.value = h
    cell.font = font("FFFFFF", 11, True)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="right" if c in "FGH" else "left", wrap_text=True, vertical="center")
ex.row_dimensions[5].height = 30
EXS = [
    ("Lucro não é caixa", "Na aba Premissas, mude o prazo de recebimento de 30 para 120 dias.",
     "O lucro não muda. O caixa das operações cai: um terço das vendas do ano fica com o cliente."),
    ("O fornecedor também financia", "Mude o prazo de pagamento aos fornecedores de 30 para 60 dias.",
     "O caixa sobe sem o lucro mudar: você vende antes de pagar."),
    ("Depreciação e imposto", "Mude a depreciação de 50 para 0.",
     "O lucro sobe, mas o caixa cai. Sem a depreciação, o lucro tributável aumenta e você paga mais IR."),
    ("Investimento não passa pela DRE", "Mude o investimento por ano (capex) de 60 para 200.",
     "O lucro fica igual. O caixa final cai 140 por ano e o imobilizado cresce no Balanço."),
    ("Quebre uma ligação", "Na aba Balanço, célula D13, apague o trecho +DRE!D14 da fórmula. Depois desfaça com Ctrl+Z.",
     "A aba Checagem mostra diferença igual ao lucro e a aba Ligações marca a ligação 2 como Quebrada."),
]
for i, (name, todo, see) in enumerate(EXS):
    r = 6 + i
    ex[f"B{r}"] = i + 1
    ex[f"C{r}"] = name
    ex[f"D{r}"] = todo
    ex[f"E{r}"] = see
    ex[f"F{r}"] = "=DRE!D14"
    ex[f"G{r}"] = "='Fluxo de Caixa'!D10"
    ex[f"H{r}"] = "=Balanço!H6"
    ex[f"B{r}"].font = font(INK2)
    ex[f"C{r}"].font = font(INK, 11, True)
    ex[f"D{r}"].font = font(INK)
    ex[f"E{r}"].font = font(INK2)
    for c in "FGH":
        ex[f"{c}{r}"].font = font(GREEN)
        ex[f"{c}{r}"].number_format = NUM
    for c in "BCDEFGH":
        ex[f"{c}{r}"].alignment = Alignment(wrap_text=c in "CDE", vertical="center", horizontal="right" if c in "FGH" else "left")
        ex[f"{c}{r}"].border = Border(bottom=thin)
    ex.row_dimensions[r].height = 48
ex["C12"] = "Valores originais"
ex["C12"].font = font(NAVY, 11, True)
ex["D12"] = "Prazo de clientes 30 · Prazo de fornecedores 30 · Depreciação 50 · Capex 60"
ex["D12"].font = font(INK2, 10)
ex["C13"] = "Resultado do modelo agora"
ex["C13"].font = font(NAVY, 11, True)
ex["D13"] = "=Checagem!C11"
ex["D13"].font = font(NAVY, 11, True)

# ---------------------------------------------------------------- Comece aqui (conteúdo)
s = start
s["B2"] = "Modelo integrado: DRE, Balanço e Fluxo de Caixa conectados"
s["B2"].font = font(NAVY, 18, True)
s["B3"] = "Saber montar cada demonstrativo separado é a parte fácil. Aqui você vê as 8 ligações que fazem os três virarem um modelo só, e fecharem."
s["B3"].font = font(INK2, 10)
rows = []
r = 5


def put(b, c, kind="text"):
    global r
    styles = {
        "h": (font(NAVY, 13, True), None),
        "step": (font(NAVY, 11, True), font(INK)),
        "bullet": (font(INK2), font(INK)),
        "note": (None, font(INK2, 10)),
        "text": (None, font(INK)),
    }
    fb, fc_ = styles[kind]
    if kind == "h":
        s[f"B{r}"] = b
        s[f"B{r}"].font = fb
    else:
        if b is not None:
            s[f"B{r}"] = b
            s[f"B{r}"].font = fb
        s[f"C{r}"] = c
        s[f"C{r}"].font = fc_
        s[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    r += 1


put("Como usar", None, "h")
put("1", "Vá na aba Premissas e mude só as células azuis (receita, crescimento, custos, prazos...).", "step")
put("2", "Acompanhe a DRE, o Balanço e o Fluxo de Caixa. Nenhuma outra célula precisa ser digitada.", "step")
put("3", "Abra a aba Ligações: ela mostra de onde sai e onde chega cada número que conecta os demonstrativos.", "step")
put("4", "Faça os 5 experimentos da aba Exercícios e confira na aba Checagem que o modelo continua fechando.", "step")
r += 1
put("As abas", None, "h")
for name, desc in [
    ("Premissas", "Os únicos números digitados do modelo."),
    ("DRE, Balanço, Fluxo de Caixa", "Os três demonstrativos de 5 anos, ligados por fórmulas."),
    ("Ligações", "As 8 conexões entre os demonstrativos, conferidas ao vivo, com o porquê de cada uma."),
    ("Checagem", "Ativo = Passivo + PL, caixa e lucro conferidos em todos os anos."),
    ("Lucro x Caixa", "A ponte entre lucro e caixa, com gráfico e indicadores de prazo e margem."),
    ("Exercícios", "Cinco testes guiados, inclusive quebrar uma ligação de propósito."),
]:
    put("•", f"{name}: {desc}", "bullet")
r += 1
put("Cores das células", None, "h")
for color, text in [(BLUE, "Azul: premissa. É o único lugar onde você digita."),
                    ("000000", "Preto: fórmula calculada na própria aba."),
                    (GREEN, "Verde: número que vem de outra aba. É aqui que os demonstrativos se conectam.")]:
    s[f"B{r}"] = "■"
    s[f"B{r}"].font = font(color, 13)
    s[f"C{r}"] = text
    s[f"C{r}"].font = font(INK)
    r += 1
r += 1
put("O que este modelo deixa de fora, de propósito", None, "h")
for t in ["Dívida e juros (e o efeito deles no caixa e no lucro)", "Dividendos e distribuição de lucro",
          "Estoques e o ciclo de caixa completo", "Cenários e análise de sensibilidade"]:
    put("•", t, "bullet")
put(None, "Antes dessas peças vem a base: no curso Fundamentos da Modelagem Financeira você monta do zero os três demonstrativos ligados, com checagem. Dívida e cenários você encontra prontos no Template Pro, que vem de bônus no curso.", "text")
r += 1
s[f"C{r}"] = "Conheça o curso Fundamentos da Modelagem Financeira →"
s[f"C{r}"].font = Font(name="Calibri", size=12, bold=True, color=AMBER_TXT, underline="single")
s[f"C{r}"].hyperlink = URL
r += 1
s[f"C{r}"] = f"Ou veja o que o curso inclui na aba {CURSO} →"
s[f"C{r}"].hyperlink = Hyperlink(ref=f"C{r}", location=f"'{CURSO}'!A1")
s[f"C{r}"].font = Font(name="Calibri", size=11, color=AMBER_TXT, underline="single")
r += 2
put(None, "Valores em R$ mil. Empresa-exemplo de 5 anos, sem dívida, para focar na ligação entre os demonstrativos.", "note")
put(None, "Modelagem Financeira na Prática · mfnapratica.com.br", "note")

# ---------------------------------------------------------------- Curso Fundamentos (oferta e links)
# Textos e preço iguais aos de src/components/fundamentos/content.ts: nada de promessa que não esteja lá.
cur_ws = sheet(wb, CURSO, "F59E0B", {"A": 3, "B": 4, "C": 100})
title(cur_ws, "Curso Fundamentos da Modelagem Financeira",
      "Você viu as ligações funcionando nesta planilha. No curso, você monta esse modelo do zero, sabendo de onde vem cada número.")
s, r = cur_ws, 5
put("Para quem trava em ligar os três demonstrativos", None, "h")
for t in ["O Balanço não fecha e você passa horas caçando a diferença entre ativo e passivo.",
          "A empresa lucrou no ano, faltou dinheiro no banco e você não consegue explicar o porquê.",
          "Você começou um curso longo, parou na metade e continua sem ligar os demonstrativos."]:
    put("•", t, "bullet")
r += 1
put("O que você monta, aula por aula", None, "h")
for i, (t, out) in enumerate([
    ("A lógica central do modelo", "o mapa de como as abas se ligam"),
    ("DRE: a estrutura do resultado", "a DRE montada no modelo"),
    ("Balanço: a fotografia", "o Balanço estruturado"),
    ("DFC: do lucro ao caixa", "o Fluxo de Caixa pelo método indireto"),
    ("Capital de giro", "o capital de giro ligado ao caixa"),
    ("As conexões que integram tudo", "os três demonstrativos conectados"),
    ("Montagem e validação", "um modelo que fecha, com checagem automática"),
], start=1):
    put(str(i), f"{t}. Você sai com {out}.", "step")
r += 1
put("O que está incluído", None, "h")
for t in ["Módulo das 3 demonstrações conectadas", "Módulos complementares com novas aulas",
          "Template do modelo integrado em Excel", "Bônus: Template Pro (vendido por R$ 97)",
          "Bônus: Starter Kit Financeiro (vendido por R$ 67,90)", "Bônus: Planilha de prompts de IA",
          "Certificado de conclusão", "Acesso vitalício"]:
    put("✓", t, "bullet")
    s[f"B{r - 1}"].font = font(GREEN, 11, True)
r += 1
put("Investimento", None, "h")
s[f"C{r}"] = "12x de R$ 19,67 no cartão, ou R$ 197 à vista"
s[f"C{r}"].font = font(NAVY, 16, True)
r += 1
put(None, "Garantia de 7 dias: se não for para você, peça o reembolso na Hotmart e receba 100% de volta.", "text")
put(None, "Pré-requisito: noção do que são DRE, Balanço e Fluxo de Caixa, e Excel básico.", "note")
r += 1
for label, link in [("Ver a página do curso →", CURSO_URL), ("Comprar agora →", CHECKOUT_URL)]:
    c = s[f"C{r}"]
    c.value = label
    c.hyperlink = link
    c.font = Font(name="Calibri", size=13, bold=True, color=INK)
    c.fill = PatternFill("solid", fgColor="F59E0B")
    c.alignment = Alignment(vertical="center", indent=1)
    s.row_dimensions[r].height = 30
    r += 1
    put(None, link, "note")  # o endereço visível, para quem abre a planilha num app que não clica em links
    r += 1

# ---------------------------------------------------------------- Formatação condicional (status)
OK_FONT, OK_FILL = Font(color=GREEN, bold=True), PatternFill("solid", bgColor="DCFCE7", fgColor="DCFCE7")
BAD_FONT, BAD_FILL = Font(color="B91C1C", bold=True), PatternFill("solid", bgColor="FEE2E2", fgColor="FEE2E2")
chk.conditional_formatting.add("C9:H9", FormulaRule(formula=['LEFT(C9,1)="F"'], font=OK_FONT, fill=OK_FILL))
chk.conditional_formatting.add("C9:H9", FormulaRule(formula=['LEFT(C9,1)="N"'], font=BAD_FONT, fill=BAD_FILL))
chk.conditional_formatting.add("C11:H11", FormulaRule(formula=['LEFT($C$11,1)="✓"'], font=OK_FONT, fill=OK_FILL))
chk.conditional_formatting.add("C11:H11", FormulaRule(formula=['LEFT($C$11,1)="✗"'], font=BAD_FONT, fill=BAD_FILL))
lig.conditional_formatting.add(f"I8:I{last}", FormulaRule(formula=['I8="Ligada"'], font=OK_FONT, fill=OK_FILL))
lig.conditional_formatting.add(f"I8:I{last}", FormulaRule(formula=['I8="Quebrada"'], font=BAD_FONT, fill=BAD_FILL))
ex.conditional_formatting.add("D13", FormulaRule(formula=['LEFT($D$13,1)="✓"'], font=OK_FONT, fill=OK_FILL))
ex.conditional_formatting.add("D13", FormulaRule(formula=['LEFT($D$13,1)="✗"'], font=BAD_FONT, fill=BAD_FILL))

# Ordem das abas: guia, premissas, demonstrativos, ligações, checagem, análise, exercícios
order = ["Comece aqui", "Premissas", "DRE", "Balanço", "Fluxo de Caixa", "Ligações", "Checagem", "Lucro x Caixa", "Exercícios", CURSO]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.calculation.fullCalcOnLoad = True  # Excel e Google Sheets calculam tudo ao abrir
wb.save(OUT)
print("ok", OUT)
